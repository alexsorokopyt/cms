import getpass
import datetime
import logging
import os
import traceback
import zipfile
from io import BytesIO

import openpyxl
import pandas as pd
import requests
from sqlalchemy import create_engine, exc, inspect


def downloadArchive(URL):
    '''
        Downloads archive from passed URL
        and extracts all its contents to download folder (specified in Parameters file)
    '''
    logging.info('Downloading started')
    filename = URL.split('/')[-1]
    req = requests.get(URL)
    logging.info(f'{filename} has been downloaded successfully')

    archive = zipfile.ZipFile(BytesIO(req.content))
    archive.extractall(result_directory)
    logging.info(f'All files from {filename} have been extracted to {DOWNLOAD_FOLDER}')


def getDataFromTable(wb, sheetName, tableName):
    '''
        Finds specified table in passed workbook
        and creates a dataframe from it
    '''
    ws = wb[sheetName]
    table = ws.tables[tableName]
    
    data = ws[table.ref]
    rows_list = []

    for row in data:
        cols = []
        for col in row:
            cols.append(col.value)
        rows_list.append(cols)

    return pd.DataFrame(data=rows_list[1:], index=None, columns=rows_list[0])


def getParameterValue(df, paramName):
    '''
        Filters Parameter table (as a dataframe) to find the value of a specific parameter
    '''
    return (df[df['Parameter'] == paramName].reset_index())['Value'][0]


current_filename = os.path.basename(__file__).rsplit('.', 1)[0]
log_file_name = current_filename + datetime.datetime.now().strftime(" %m.%d.%Y %H.%M.%S") + '.log'
logging.basicConfig(filename = log_file_name, level = logging.INFO, format = '%(asctime)s:%(levelname)s:%(message)s')

PARAMETERS_PATH = 'C:\\Users\\user\\Documents\\cms\\Parameters.xlsx'
ALPHA_TABLE = 'rpt_alpha'
RPT_TABLE = 'rpt'
NMRC_TABLE = 'rpt_nmrc'

wb = openpyxl.load_workbook(PARAMETERS_PATH)

parameters = getDataFromTable(wb, 'Parameters', 'Parameters')

HOST = getParameterValue(parameters, 'Host')
DB = getParameterValue(parameters, 'Database')
DOWNLOAD_FOLDER = getParameterValue(parameters, 'Download folder')
DOWNLOAD_FOLDER_TYPE = getParameterValue(parameters, 'Download folder path type')
SHOULD_CLEAR = getParameterValue(parameters, 'Clear download folder')
TABLES_BEHAVIOUR = getParameterValue(parameters, 'If tables exist')

URLs = getDataFromTable(wb, 'Parameters', 'URLs')['URL']
mappings = getDataFromTable(wb, 'Parameters', 'Mappings')
mappings_tables = list(mappings['DB table'])
mappings_parts = list(mappings['Part of file name'])
mappings_defHeaders = list(mappings['Default headers'])
mappings_headers = list(mappings['DB columns'])

wb.close()

logging.info(f'''
PASSED PARAMETERS:

Host: {HOST}
DB: {DB}
Downloads folder: {DOWNLOAD_FOLDER}
Clear download folder: {SHOULD_CLEAR}
Behaviour if tables exist: {TABLES_BEHAVIOUR}

URLS: 
{URLs}

MAPPINGS: 
{mappings}''')

current_directory = os.getcwd()
logging.info('Current working directory: ' + current_directory)

result_directory = ( current_directory + '\\' if DOWNLOAD_FOLDER_TYPE == 'relative' else '' ) + DOWNLOAD_FOLDER

if SHOULD_CLEAR == 'yes':
    if os.path.exists(result_directory):
        for file in os.listdir(result_directory):
            os.remove(result_directory + '\\' + file)
            logging.info('File has been removed: ' + file)

        os.rmdir(result_directory)
        logging.info('Downloads folder has been removed')
    else:
        logging.info('Downloads folder doesn\'t exist')

for url in URLs:
    downloadArchive(url)

try:
    DB_USER = input('Enter database username: ')
    DB_PASSWORD = getpass.getpass('Enter database password: ')

    db_engine = create_engine(f'mysql://{DB_USER}:{DB_PASSWORD}@{HOST}/{DB}')

    with db_engine.connect() as connection:
        logging.info('MySQL engine has been successfully created')

        if TABLES_BEHAVIOUR == 'replace':
            # Getting all shema tables names
            inspector = inspect(db_engine)
            table_names = inspector.get_table_names()

            trans = connection.begin()
            connection.execute('SET FOREIGN_KEY_CHECKS = 0;')
            for table in table_names:
                # Sequentially truncate each existing table
                connection.execute(f'TRUNCATE {table}')
                logging.info(f'Table {table} has been truncated')
            connection.execute('SET FOREIGN_KEY_CHECKS = 1;')
            trans.commit()

        # Iterrating over all files in Downloads folder
        for file in os.listdir(result_directory):
            logging.info('Started loading data from file ' + result_directory + '\\' + file)

            # Gettings current file extension
            file_extension = os.path.splitext(file)[-1]
            
            if file_extension.lower() != '.csv':
                logging.error(f'Unsupported file extension: {file}')
            else:
                # Iterrating over mappings table
                for index, path_part in enumerate(mappings_parts):
                    # If part of file name was found in current file name
                    # Then get 'table name' and 'headers' properties
                    if path_part.lower() in file.lower():

                        defHeaders = mappings_defHeaders[index]
                        headerManual = mappings_headers[index]
                        table_name = mappings_tables[index]

                        # If there are no headers in the file and no headers in the Parameters file, then skip this file
                        if defHeaders == 'no':
                            if headerManual == None:
                                logging.critical(f'No specified headers for file {file}')
                                break
                            else:
                                header = headerManual.split(';')
                                df = pd.read_csv(result_directory + '\\' + file, 
                                                names=header)
                        else:
                            header = [0]
                            df = pd.read_csv(result_directory + '\\' + file, 
                                            header=header)
                  
                        # 'append' creates table if it doesn't exist
                        df.to_sql(con=db_engine, 
                                  schema=DB, 
                                  name=table_name, 
                                  if_exists='append', 
                                  index=False,
                                  chunksize=1000)

                        logging.critical(f'{len(df.index)} new rows were successfully loaded to table {table_name} from {file}')
                        break

        connection.close()

except exc.SQLAlchemyError as e:
    logging.error(e)
except Exception as e:
    logging.error(traceback.format_exc())
finally:
    db_engine.dispose()
    logging.info(f'MySQL engine has been disposed')
